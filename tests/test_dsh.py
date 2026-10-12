"""Offline regression tests; all writes stay inside temporary directories."""
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import install_dsh as installer
import dsh_runtime as runtime
import research_quality_gate as gate
import check_skill_integrity as integrity


class Workspace(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="论文 空格-")
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.home = self.base / "DSH 配置"
        self.paper = self.base / "论文 工作区"
        self.paper.mkdir()

    def write(self, rel, text):
        path = self.paper / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")

    def manuscript(self, role="body", filename="06_结论.md", citations=""):
        self.write("plan/chapter-architecture.md",
                   f"- chapters/{filename} | role={role} | min_chars=10 | agent=optional | placeholders=no\n")
        prose = "这是用于离线门禁测试的示例段落，不是真实论文，也不代表实际研究结论。" * 4
        self.write(f"chapters/{filename}", "\n\n".join([prose] * 6) + citations)
        self.write("refs/evidence-map.md", "测试资料映射，未调用外部模型。")


class InstallationTests(Workspace):
    def test_install_repeat_and_uninstall(self):
        installer.install(ROOT, self.home)
        installer.install(ROOT, self.home)
        links = list((self.home / "skills").iterdir())
        self.assertEqual(len(links), 19)
        for link in links:
            self.assertTrue(link.is_symlink())
            self.assertEqual(runtime.resolve_root(link), ROOT)
        foreign = self.home / "skills/other-skill"
        foreign.mkdir()
        installer.uninstall(self.home)
        self.assertEqual(list((self.home / "skills").iterdir()), [foreign])

    def test_conflict_preflight_does_not_partially_install(self):
        conflict = self.home / "skills/writing-humanities"
        conflict.mkdir(parents=True)
        with self.assertRaises(installer.InstallError):
            installer.install(ROOT, self.home)
        self.assertEqual(list((self.home / "skills").iterdir()), [conflict])

    def test_unowned_identical_link_is_a_conflict(self):
        link = self.home / "skills/writing-humanities"
        link.parent.mkdir(parents=True)
        link.symlink_to(ROOT / "skills/writing-humanities")
        with self.assertRaises(installer.InstallError):
            installer.install(ROOT, self.home)

    def test_changed_managed_link_is_not_removed(self):
        installer.install(ROOT, self.home)
        link = self.home / "skills/writing-humanities"
        link.unlink()
        link.symlink_to(self.paper)
        with self.assertRaises(installer.InstallError):
            installer.uninstall(self.home)
        self.assertEqual(link.resolve(), self.paper)
        self.assertEqual(len(list((self.home / "skills").iterdir())), 19)

    def test_source_relocation_updates_owned_links(self):
        installer.install(ROOT, self.home)
        moved = self.base / "搬迁后的 技能"
        shutil.copytree(ROOT, moved, ignore=shutil.ignore_patterns("__pycache__", ".venv*"))
        installer.install(moved, self.home)
        self.assertEqual((self.home / "skills/writing-humanities").resolve(), moved / "skills/writing-humanities")

    def test_dsh_home_environment(self):
        with patch.dict(os.environ, {"DSH_HOME": str(self.home)}):
            self.assertEqual(installer.default_home(), self.home)

    def test_resource_resolution_from_file(self):
        self.assertEqual(runtime.resolve_root(ROOT / "skills/writing-humanities/SKILL.md"), ROOT)

    def test_invalid_resource_location(self):
        with self.assertRaises(ValueError):
            runtime.resolve_root(self.paper)

    def test_plan_initialization_from_other_cwd_preserves_files(self):
        self.write("plan/notes.md", "用户原有记录")
        subprocess.run(["bash", str(ROOT / "scripts/init_plan.sh"), str(self.paper)],
                       cwd=self.base, check=True, capture_output=True)
        self.assertEqual((self.paper / "plan/notes.md").read_text(), "用户原有记录")
        self.assertTrue((self.paper / "plan/project-overview.md").is_file())

    def test_missing_managed_link_is_repaired(self):
        installer.install(ROOT, self.home)
        path = self.home / "skills/writing-humanities"
        path.unlink()
        installer.install(ROOT, self.home)
        self.assertEqual(runtime.resolve_root(path), ROOT)

    def test_uninstall_without_install_is_idempotent(self):
        self.assertEqual(installer.uninstall(self.home), {})

    def test_invalid_manifest_cannot_remove_other_files(self):
        path = installer.manifest_path(self.home)
        path.parent.mkdir(parents=True)
        path.write_text(json.dumps({"version": 1, "links": {"../other": str(self.paper)}}))
        with self.assertRaises(installer.InstallError):
            installer.uninstall(self.home)
        self.assertTrue(self.paper.is_dir())

    def test_manifest_write_failure_rolls_back_links(self):
        with patch.object(installer, "write_manifest", side_effect=OSError("fixture disk failure")):
            with self.assertRaises(OSError):
                installer.install(ROOT, self.home)
        self.assertEqual(list((self.home / "skills").iterdir()), [])

    def test_all_skills_have_runtime_contract(self):
        self.assertEqual(integrity.check(ROOT), [])


class QualityTests(Workspace):
    def test_characterization_old_numeric_citation_bug(self):
        # Frozen original regex: it recognized [FL-05], but missed undergraduate numeric citations.
        import re
        old = r"\uFF08[^\uFF09]*(19|20)\d{2}[^\uFF09]*\uFF09|\([A-Z][^)]*(19|20)\d{2}[^)]*\)|\[[A-Za-z]+-?\d+\]"
        self.assertEqual(len(re.findall(old, "[1] [1,2] [1-3]")), 0)

    def test_numeric_citation_forms(self):
        for value in ("[1]", "[1,2]", "[1-3]", "[1, 2-4]", "[FL-05]", "（张三，2024）", "(Smith, 2024)"):
            with self.subTest(value=value):
                self.assertEqual(gate.count_citations(value), 1)

    def test_invalid_numeric_forms(self):
        for value in ("[0]", "[01]", "[1,]", "[x]", "[1,0]"):
            self.assertEqual(gate.count_citations(value), 0, value)

    def test_chinese_intro_role_accepts_numeric_citations(self):
        self.manuscript("introduction", "第一章 绪论.md", " [1]" * 5)
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_intro_without_citations_fails(self):
        self.manuscript("introduction")
        self.assertTrue(any("citation" in x for x in gate.check_project(self.paper)[0]))

    def test_related_work_requires_eight_markers(self):
        self.manuscript("related-work", "文献综述.md", " [1]" * 7)
        self.assertTrue(any("citation" in x for x in gate.check_project(self.paper)[0]))

    def test_conclusion_number_does_not_imply_experiment(self):
        self.manuscript("conclusion")
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_empirical_role_requires_artifacts(self):
        self.manuscript("empirical-results", "调查发现.md")
        self.assertTrue(any("experiment-protocol" in x for x in gate.check_project(self.paper)[0]))

    def test_missing_evidence_fails(self):
        self.manuscript()
        (self.paper / "refs/evidence-map.md").unlink()
        self.assertTrue(any("evidence map" in x for x in gate.check_project(self.paper)[0]))

    def test_submission_placeholders_fail(self):
        self.manuscript()
        self.write("plan/notes.md", "TODO")
        self.assertEqual(gate.check_project(self.paper)[0], [])
        self.assertTrue(gate.check_project(self.paper, submission=True)[0])

    def test_unresolved_chapter_placeholders_fail(self):
        self.manuscript()
        with (self.paper / "chapters/06_结论.md").open("a") as f:
            f.write("TODO")
        self.assertTrue(any("placeholder" in x for x in gate.check_project(self.paper)[0]))

    def test_missing_role_fails_instead_of_guessing(self):
        self.manuscript()
        self.write("plan/chapter-architecture.md", "- chapters/06_结论.md | min_chars=10\n")
        self.assertTrue(any("role" in x for x in gate.check_project(self.paper)[0]))

    def test_missing_architecture_fails(self):
        self.manuscript()
        (self.paper / "plan/chapter-architecture.md").unlink()
        self.assertTrue(any("chapter-architecture" in x for x in gate.check_project(self.paper)[0]))

    def test_unlisted_and_missing_chapters_fail(self):
        self.manuscript()
        self.write("chapters/额外.md", "正文")
        (self.paper / "chapters/06_结论.md").unlink()
        failures = gate.check_project(self.paper)[0]
        self.assertTrue(any("Unexpected" in x for x in failures))
        self.assertTrue(any("missing file" in x for x in failures))

    def test_required_provenance(self):
        self.manuscript()
        path = self.paper / "plan/chapter-architecture.md"
        path.write_text(path.read_text().replace("agent=optional", "agent=required"))
        self.assertTrue(any("provenance" in x for x in gate.check_project(self.paper)[0]))
        self.write("plan/chapter-agent-provenance.md", "chapters/06_结论.md | status=ACCEPTED")
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_path_traversal_is_rejected(self):
        self.manuscript()
        self.write("plan/chapter-architecture.md", "- chapters/../../outside.md | role=body\n")
        self.assertTrue(any("path" in x.lower() for x in gate.check_project(self.paper)[0]))

    def test_symlink_outside_project_is_rejected(self):
        self.manuscript()
        chapter = self.paper / "chapters/06_结论.md"
        chapter.unlink()
        external = self.base / "外部.md"
        external.write_text("外部文件")
        chapter.symlink_to(external)
        self.assertTrue(any("path" in x.lower() for x in gate.check_project(self.paper)[0]))

    def test_role_overrides_numeric_filename_in_characterization(self):
        # Original gate inferred results solely from the prefix 06_.
        import re
        self.assertTrue(re.search(r"Results|results|06_", "06_结论.md"))
        self.manuscript("conclusion", "06_结论.md")
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_submission_forbids_mock_planning_files(self):
        self.manuscript()
        self.write("tables/mock_fixture.csv", "PLANNING DATA - replace before submission\nx,y\n1,2")
        self.assertEqual(gate.check_project(self.paper)[0], [])
        self.assertTrue(any("mock/synthetic" in x for x in gate.check_project(self.paper, True)[0]))

    def test_role_threshold_can_follow_school_requirements(self):
        self.manuscript("introduction")
        path = self.paper / "plan/chapter-architecture.md"
        path.write_text(path.read_text().strip() + " | min_citations=0 | min_paragraphs=2\n")
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_reference_role_is_exempt_from_prose_threshold(self):
        self.manuscript("references", "参考文献.md")
        self.write("chapters/参考文献.md", "[1] 示例来源文献信息，来源真实性另行核对。")
        self.assertEqual(gate.check_project(self.paper)[0], [])

    def test_quality_cli_returns_failure_code(self):
        self.manuscript("introduction")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/research_quality_gate.py"), "--project-path", str(self.paper)],
                                capture_output=True, text=True, cwd=self.base)
        self.assertEqual(result.returncode, 1)
        self.assertIn("citation", result.stdout)


class DependencyTests(Workspace):
    def test_pdf_cli_offline(self):
        import fitz
        pdf = self.paper / "文献 示例.pdf"
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), "Offline literature fixture")
            doc.save(pdf)
        out = self.paper / "文本.txt"
        subprocess.run([sys.executable, str(ROOT / "scripts/pdf_parser.py"), str(pdf), "--output", str(out)],
                       cwd=self.paper, check=True, capture_output=True)
        self.assertIn("Offline literature fixture", out.read_text())

    def test_science_and_plot(self):
        os.environ["MPLBACKEND"] = "Agg"
        import numpy as np
        import pandas as pd
        import scipy.stats as stats
        import matplotlib.pyplot as plt
        import seaborn
        import sklearn
        import statsmodels.api as sm
        import pingouin
        import openpyxl
        import plotly
        data = pd.DataFrame({"value": np.arange(8)})
        self.assertGreater(stats.ttest_1samp(data.value, 0).statistic, 0)
        self.assertIn("T", pingouin.ttest(np.arange(8), 0).columns)
        self.assertTrue(np.isfinite(sm.OLS(data.value, sm.add_constant(np.arange(8))).fit().params).all())
        fig, ax = plt.subplots()
        ax.plot(data.value)
        fig.savefig(self.paper / "figure.png")
        plt.close(fig)
        self.assertGreater((self.paper / "figure.png").stat().st_size, 0)

    def test_chinese_figure_template_offline(self):
        os.environ["MPLBACKEND"] = "Agg"
        import matplotlib.pyplot as plt
        import numpy as np
        spec = importlib.util.spec_from_file_location("figure_template", ROOT / "templates/figure-template.py")
        template = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(template)
        template.setup_plot_style()
        fig, ax = plt.subplots()
        template.plot_line_chart(ax, np.arange(4), {"示例数据": np.arange(4)}, xlabel="时间", ylabel="数值")
        template.save_figure(fig, self.paper, "中文 图表")
        plt.close(fig)
        self.assertTrue((self.paper / "中文 图表.svg").is_file())
        self.assertTrue((self.paper / "中文 图表.png").is_file())

    def test_scholar_crossref_mock(self):
        import scholar_search
        response = Mock()
        response.status_code = 200
        response.json.return_value = {"message": {"items": [{
            "title": ["Offline fixture"], "author": [{"given": "A", "family": "B"}],
            "published-print": {"date-parts": [[2024]]}, "DOI": "10.0000/fixture"
        }]}}
        search = scholar_search.ScholarSearch()
        with patch.object(search.session, "get", return_value=response) as get:
            papers = search._search_crossref("fixture", 2020, 2026, 1)
        self.assertTrue(get.called)
        self.assertEqual(papers[0]["title"], "Offline fixture")


if __name__ == "__main__":
    unittest.main()
